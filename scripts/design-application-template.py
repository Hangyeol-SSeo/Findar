"""Author the finance template from a retained working copy (design-time only).

Usage: python scripts/design-application-template.py SOURCE.docx OUTPUT.docx
Requires python-docx only for authoring; the production filler uses stdlib.
"""
from pathlib import Path
import sys
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor
from docx.table import _Cell

FONT = 'Nanum Gothic'
WIDTH = 174
FILL = 'F0F1F2'
BORDER = 'B9BEC4'


def prop(parent, name, **attrs):
    node = parent.find(qn('w:' + name))
    if node is None:
        node = OxmlElement('w:' + name)
        parent.append(node)
    for key, value in attrs.items():
        node.set(qn('w:' + key), str(value))
    return node


def font(run, size=10.5, bold=False, color='000000'):
    run.font.name, run.font.size, run.font.bold = FONT, Pt(size), bold
    run.font.color.rgb = RGBColor.from_string(color)
    rf = run._r.get_or_add_rPr().get_or_add_rFonts()
    for key in ('ascii', 'hAnsi', 'eastAsia', 'cs'):
        rf.set(qn('w:' + key), FONT)
    return run


def mark(node, key, number):
    value = 'findar_' + key.replace('.', '_')
    if node.tag == qn('w:tbl'):
        prop(node.tblPr, 'tblCaption', val=value)
    else:
        start = OxmlElement('w:bookmarkStart')
        start.set(qn('w:id'), str(number))
        start.set(qn('w:name'), value)
        node.insert(1 if node.find(qn('w:pPr')) is not None else 0, start)
        end = OxmlElement('w:bookmarkEnd')
        end.set(qn('w:id'), str(number))
        node.append(end)


def build(source_path, output_path):
    if Path(source_path).resolve() == Path(output_path).resolve():
        raise ValueError('Retain the source; write to another path.')
    doc = Document(source_path)
    for child in list(doc._element.body):
        if child.tag != qn('w:sectPr'):
            doc._element.body.remove(child)
    section = doc.sections[0]
    section.page_width, section.page_height = Mm(210), Mm(297)
    section.left_margin = section.right_margin = Mm(18)
    section.top_margin = section.bottom_margin = Mm(17)
    section.header_distance = section.footer_distance = Mm(8)
    section.different_first_page_header_footer = False
    for name, size in [('Normal', 10.5), ('Title', 20), ('Heading 1', 11), ('Heading 2', 10.5), ('Subtitle', 9.5), ('Header', 8), ('Footer', 8)]:
        s = doc.styles[name]
        s.font.name, s.font.size, s.font.color.rgb = FONT, Pt(size), RGBColor(0, 0, 0)
        prop(s.element.get_or_add_rPr(), 'rFonts', ascii=FONT, hAnsi=FONT, eastAsia=FONT, cs=FONT)
        pp = s.element.get_or_add_pPr()
        for old in list(pp):
            pp.remove(old)
        prop(pp, 'snapToGrid', val=0)
        pf = s.paragraph_format
        pf.space_before, pf.space_after, pf.line_spacing = Pt(0), Pt(4), 1.2
        pf.widow_control = True
        s.font.bold = name in ('Title', 'Heading 1', 'Heading 2')
        if s.font.bold:
            pf.keep_with_next = True
        if name == 'Title':
            pf.alignment, pf.space_after = WD_ALIGN_PARAGRAPH.CENTER, Pt(12)
        elif name == 'Heading 1':
            pf.space_before, pf.space_after = Pt(8), Pt(4)
        elif name == 'Heading 2':
            pf.space_before, pf.space_after = Pt(7), Pt(4)
    for part in (section.header, section.footer):
        for child in list(part._element):
            part._element.remove(child)
        part.add_paragraph()
    footer = section.footer.paragraphs[0]
    footer.style = doc.styles['Footer']
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    font(footer.add_run('[성명]  |  '), 8)
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), 'PAGE')
    footer._p.append(field)
    doc.core_properties.title = '입사지원서'
    doc.core_properties.author = ''
    doc.core_properties.last_modified_by = ''
    counter = [100]

    def p(key, value, style='Normal', size=None, after=None, page_break=False, label=None):
        para = doc.add_paragraph(style=style)
        if label:
            font(para.add_run(label + '  '), size or 10.5, True)
        font(para.add_run(value), size or (20 if style == 'Title' else 11 if style == 'Heading 1' else 10.5), style.startswith('Heading') or style == 'Title')
        if after is not None:
            para.paragraph_format.space_after = Pt(after)
        if page_break:
            para.paragraph_format.page_break_before = True
        mark(para._p, key, counter[0]); counter[0] += 1
        return para

    def table(key, widths, rows, header=False, labels=False, size=10):
        t = doc.add_table(rows=0, cols=len(widths))
        t.autofit, t.alignment = False, WD_TABLE_ALIGNMENT.CENTER
        for col, width in zip(t.columns, widths):
            col.width = Mm(width)
        tp = t._tbl.tblPr
        prop(tp, 'tblW', w=round(WIDTH / 25.4 * 1440), type='dxa')
        prop(tp, 'tblInd', w=0, type='dxa')
        borders = prop(tp, 'tblBorders')
        for side in ('top', 'bottom', 'left', 'right', 'insideH', 'insideV'):
            prop(borders, side, val='single', sz=5, color=BORDER)
        margins = prop(tp, 'tblCellMar')
        for side, width in [('top', 60), ('bottom', 60), ('left', 100), ('right', 100)]:
            prop(margins, side, w=width, type='dxa')
        for ri, values in enumerate(rows):
            row = t.add_row()
            prop(row._tr.get_or_add_trPr(), 'cantSplit')
            if header and ri == 0:
                prop(row._tr.get_or_add_trPr(), 'tblHeader')
            if len(values) == 2 and len(widths) == 4:
                row.cells[1].merge(row.cells[3])
            for ci, (tc, value) in enumerate(zip(row._tr.tc_lst, values)):
                cell = _Cell(tc, t)
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                para = cell.paragraphs[0]
                para.paragraph_format.space_after = Pt(0)
                para.paragraph_format.line_spacing = 1.15
                para.paragraph_format.keep_with_next = False
                prop(para._p.get_or_add_pPr(), 'snapToGrid', val=0)
                is_label, is_header = labels and ci % 2 == 0, header and ri == 0
                centered_value = not header or ri == 0
                if header and ri > 0:
                    centered_value = (ci == 0 or (key == 'education.table' and ci >= 2)
                                      or (key == 'work.table' and ci == 2)
                                      or (key == 'credentials.table' and ci >= 3))
                para.alignment = WD_ALIGN_PARAGRAPH.CENTER if is_header or is_label or (header and centered_value) else WD_ALIGN_PARAGRAPH.LEFT
                font(para.add_run(value), size, is_label or is_header)
                if is_label or is_header:
                    prop(cell._tc.get_or_add_tcPr(), 'shd', fill=FILL, val='clear')
        mark(t._tbl, key, counter[0]); counter[0] += 1
        return t

    p('resume.title', '입사지원서', 'Title')
    target = table('resume.target', [24, 80, 24, 46], [
        ['지원 회사', '[지원 회사]'],
        ['지원 부서', '[지원 팀 또는 직무]', '지원 구분', '[신입 / 경력 / 인턴]'],
    ], labels=True, size=11)
    for row in target.rows:
        for para in row.cells[1].paragraphs:
            for run in para.runs:
                run.bold = True
    p('personal.heading', '1. 기본 인적사항', 'Heading 1')
    table('personal.table', [24, 63, 24, 63], [
        ['성명', '[성명]', '연락처', '[연락처]'], ['이메일', '[이메일]'],
    ], labels=True, size=10.5)
    p('education.heading', '2. 학력사항', 'Heading 1')
    table('education.table', [31, 68, 48, 27], [
        ['재학 기간', '학교 및 전공', '학위 및 졸업 구분', '학점'],
        ['[입학] - [졸업]', '[학교명]\n[전공명]', '[학위 / 졸업 구분]', '[취득 / 만점]'],
    ], header=True)
    p('work.heading', '3. 경력사항', 'Heading 1')
    table('work.table', [31, 48, 35, 60], [
        ['근무 기간', '회사 및 소속', '직위 및 고용 형태', '담당업무'],
        ['[입사] - [퇴사]', '[회사명]\n[부서명]', '[직위]\n[고용 형태]', '[담당업무]'],
    ], header=True)
    p('credentials.heading', '4. 자격 및 어학', 'Heading 1')
    table('credentials.table', [15, 49, 47, 31, 32], [
        ['구분', '자격 및 시험명', '발급 또는 시행 기관', '취득일', '등급 및 점수'],
        ['자격', '[자격명]', '[발급 기관]', '[취득일]', '[등급]'],
        ['어학', '[시험명]', '[시행 기관]', '[응시일]', '[점수]'],
    ], header=True)
    p('summary.heading', '5. 활용 역량 및 참고 자료', 'Heading 1')
    table('summary.table', [25, 149], [
        ['활용 역량', '[활용 가능한 도구 및 실제 활용 범위]'],
        ['참고 자료', '[포트폴리오 또는 확인 자료 링크]'],
    ], labels=True)
    p('details.title', '경력 및 경험 상세', 'Title', size=18, page_break=True)
    identity = p('details.subtitle', '[지원 회사]  |  [지원 팀 또는 직무]  |  [성명]', 'Subtitle', size=9.5, after=10)
    identity.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p('details.work.heading', '업무 경험', 'Heading 1')
    p('details.work.title', '[회사명 / 부서 및 직위]', 'Heading 2')
    p('details.meta', '[재직 기간 / 고용 형태]', size=9.5, after=6)
    p('details.work.line', '[직접 맡은 업무 범위와 실제 수행한 내용을 작성해주세요.]', label='담당 업무')
    p('details.work.contribution', '[본인이 판단하고 실행한 내용과 확인 가능한 결과를 작성해주세요.]', label='본인 기여 및 결과')
    p('details.project.heading', '프로젝트 및 학습 경험', 'Heading 1')
    p('details.project.title', '[프로젝트 또는 활동명]', 'Heading 2')
    p('details.project.meta', '[기간 / 실제 경험 구분]', size=9.5, after=6)
    p('details.project.line', '[본인이 수행한 범위와 실제 과정을 작성해주세요.]', label='역할 및 내용')
    p('details.project.result', '[확인한 결과와 한계를 작성해주세요.]', label='결과 및 한계')
    p('essay.title', '자기소개서', 'Title', size=18, page_break=True)
    identity = p('essay.identity', '[지원 회사]  |  [지원 팀 또는 직무]  |  [성명]', 'Subtitle', size=9.5, after=14)
    identity.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for i, (question, guide) in enumerate([
        ('지원동기', '[해당 회사와 직무에 지원한 이유, 실제 준비한 과정을 작성해주세요.]'),
        ('직무 관련 역량', '[지원 업무와 관련된 경험에서 본인이 수행한 내용과 결과를 작성해주세요.]'),
        ('일하는 방식', '[문제 해결, 검증 또는 협업에서 본인의 행동을 보여주는 경험을 작성해주세요.]'),
    ]):
        key = '' if i == 0 else '.' + str(i + 1)
        question_p = p('essay.question' + key, question, 'Heading 1', size=11)
        prop(question_p._p.get_or_add_pPr(), 'shd', val='clear', fill=FILL)
        question_p.paragraph_format.space_before = Pt(12)
        question_p.paragraph_format.space_after = Pt(8)
        body = p('essay.body' + key, guide, size=11, after=80)
        body.paragraph_format.line_spacing = 1.45
        for run in body.runs:
            run.font.color.rgb = RGBColor.from_string('72777D')
    doc.save(output_path)


if __name__ == '__main__':
    build(sys.argv[1], sys.argv[2])
